from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


SOURCE = Path(r"C:\Users\lxj\Desktop\李小嘉2026年山西省研究生实践创新项目申请书.docx")
OUT = Path(r"E:\pythonproject\3d_cnn_classify _K\tmp\innovation_application_2026\template_inspection.json")


def text_of_cell(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip()


def font_info(style) -> dict[str, object]:
    font = style.font
    return {
        "name": font.name,
        "east_asia": style.element.rPr.rFonts.get(qn("w:eastAsia")) if style.element.rPr is not None and style.element.rPr.rFonts is not None else None,
        "size_pt": font.size.pt if font.size else None,
        "bold": font.bold,
        "italic": font.italic,
    }


def main() -> None:
    doc = Document(SOURCE)
    sections = []
    for index, section in enumerate(doc.sections):
        sections.append(
            {
                "index": index,
                "page_width": section.page_width,
                "page_height": section.page_height,
                "top_margin": section.top_margin,
                "bottom_margin": section.bottom_margin,
                "left_margin": section.left_margin,
                "right_margin": section.right_margin,
                "header_distance": section.header_distance,
                "footer_distance": section.footer_distance,
                "header": [p.text for p in section.header.paragraphs],
                "footer": [p.text for p in section.footer.paragraphs],
            }
        )

    paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip():
            paragraphs.append(
                {
                    "index": index,
                    "style": paragraph.style.name,
                    "alignment": str(paragraph.alignment),
                    "text": paragraph.text,
                    "runs": [
                        {
                            "text": run.text,
                            "font_name": run.font.name,
                            "size_pt": run.font.size.pt if run.font.size else None,
                            "bold": run.bold,
                            "italic": run.italic,
                        }
                        for run in paragraph.runs
                        if run.text
                    ],
                }
            )

    tables = []
    for index, table in enumerate(doc.tables):
        cell_ids = [[id(cell._tc) for cell in row.cells] for row in table.rows]
        tables.append(
            {
                "index": index,
                "style": table.style.name if table.style else None,
                "rows": len(table.rows),
                "cols": len(table.columns),
                "cells": [[text_of_cell(cell) for cell in row.cells] for row in table.rows],
                "cell_ids": cell_ids,
                "widths": [[cell.width for cell in row.cells] for row in table.rows],
                "row_xml": [row._tr.xml for row in table.rows],
            }
        )

    styles = []
    for style in doc.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            styles.append(
                {
                    "name": style.name,
                    "based_on": style.base_style.name if style.base_style else None,
                    "paragraph": {
                        "alignment": str(style.paragraph_format.alignment),
                        "space_before": style.paragraph_format.space_before.pt if style.paragraph_format.space_before else None,
                        "space_after": style.paragraph_format.space_after.pt if style.paragraph_format.space_after else None,
                        "line_spacing": str(style.paragraph_format.line_spacing),
                        "first_line_indent": style.paragraph_format.first_line_indent,
                        "left_indent": style.paragraph_format.left_indent,
                    },
                    "font": font_info(style),
                }
            )

    result = {
        "source": str(SOURCE),
        "paragraph_count": len(doc.paragraphs),
        "nonempty_paragraphs": paragraphs,
        "table_count": len(doc.tables),
        "tables": tables,
        "inline_shape_count": len(doc.inline_shapes),
        "sections": sections,
        "paragraph_styles": styles,
    }
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
