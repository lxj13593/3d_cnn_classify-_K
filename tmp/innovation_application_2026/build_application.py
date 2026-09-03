from __future__ import annotations

import copy
import hashlib
import os
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"E:\\pythonproject\\3d_cnn_classify _K")
WORK = ROOT / "tmp" / "innovation_application_2026"
OUTPUT_DIR = ROOT / "deliverables"
SOURCE_TEMPLATE = Path(r"C:\\Users\\lxj\\Desktop") / (
    "\u674e\u5c0f\u56092026\u5e74\u5c71\u897f\u7701\u7814\u7a76\u751f\u5b9e\u8df5\u521b\u65b0\u9879\u76ee\u7533\u8bf7\u4e66.docx"
)
OUTPUT = OUTPUT_DIR / (
    "\u674e\u5c0f\u56092026\u5e74\u5c71\u897f\u7701\u7814\u7a76\u751f\u5b9e\u8df5\u521b\u65b0\u9879\u76ee\u7533\u8bf7\u4e66_"
    "\u94bb\u5b54\u4e09\u7ef4\u7f3a\u9677\u5206\u7c7b\u9879\u76ee.docx"
)
TITLE = "\u9762\u5411\u94bb\u5b54\u5fae\u5c0f\u7f3a\u9677\u68c0\u6d4b\u7684\u5934\u90e8\u4e09\u7ef4\u4f53\u79ef\u667a\u80fd\u5206\u7c7b\u5173\u952e\u6280\u672f\u53ca\u5e94\u7528"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def get_font(size: int, bold: bool = False):
    candidates = [
        r"C:\\Windows\\Fonts\\msyhbd.ttc" if bold else r"C:\\Windows\\Fonts\\msyh.ttc",
        r"C:\\Windows\\Fonts\\simhei.ttf" if bold else r"C:\\Windows\\Fonts\\simsun.ttc",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size, index=0)
    return ImageFont.load_default()


FONTS = {}


def font(size: int, bold: bool = False):
    key = (size, bold)
    if key not in FONTS:
        FONTS[key] = get_font(size, bold)
    return FONTS[key]


NAVY = "#174A72"
BLUE = "#2F77A8"
TEAL = "#2C8A86"
LIGHT_BLUE = "#E8F3FA"
LIGHT_TEAL = "#E6F5F1"
LIGHT_GRAY = "#F4F7F9"
MID_GRAY = "#617280"
TEXT = "#23313B"
ORANGE = "#E8923A"


def rounded_box(draw, box, fill, outline=BLUE, radius=24, width=3):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def centered_multiline(draw, box, text, size=34, color=TEXT, bold=False, spacing=10):
    x0, y0, x1, y1 = box
    f = font(size, bold)
    max_width = x1 - x0 - 32
    lines, current = [], ""
    for char in text:
        proposal = current + char
        if draw.textbbox((0, 0), proposal, font=f)[2] <= max_width:
            current = proposal
        else:
            if current:
                lines.append(current)
            current = char
    if current:
        lines.append(current)
    line_h = draw.textbbox((0, 0), "国", font=f)[3] + spacing
    total_h = len(lines) * line_h - spacing
    y = y0 + (y1 - y0 - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=f)
        x = x0 + (x1 - x0 - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=f, fill=color)
        y += line_h


def arrow(draw, start, end, color=BLUE, width=7):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = end
    x0, y0 = start
    if abs(x1 - x0) >= abs(y1 - y0):
        direction = 1 if x1 > x0 else -1
        draw.polygon([(x1, y1), (x1 - 20 * direction, y1 - 13), (x1 - 20 * direction, y1 + 13)], fill=color)
    else:
        direction = 1 if y1 > y0 else -1
        draw.polygon([(x1, y1), (x1 - 13, y1 - 20 * direction), (x1 + 13, y1 - 20 * direction)], fill=color)


def canvas(title: str, subtitle: str, height: int = 1180):
    img = Image.new("RGB", (2200, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, 2200, 130), fill=NAVY)
    draw.text((80, 34), title, font=font(48, True), fill="#FFFFFF")
    draw.text((82, 92), subtitle, font=font(24), fill="#D8ECF8")
    return img, draw


def draw_technical_route(path: Path):
    img, draw = canvas("钻孔三维缺陷智能分类技术路线", "以元数据先验约束数据处理，以固定划分验证泛化能力", 1420)
    cols = [120, 620, 1120, 1620]
    labels = [
        ("三维体积数据与元数据", "RAW 体数据 + TOML 标注\n22 板训练验证 / 3 板锁定外测"),
        ("物理一致的数据构建", "序列映射、头部 ROI 裁剪\nbHeadUp 方向统一、鲁棒归一化"),
        ("三维特征学习", "Head40 3D ResNet18\n多尺度核比较、轻量增强、五折训练"),
        ("可信评估与应用", "OOF 验证、锁定外部板测试\n阈值分析、可解释性、辅助质检"),
    ]
    for x, (head, body) in zip(cols, labels):
        rounded_box(draw, (x, 300, x + 360, 665), LIGHT_BLUE if x in (120,1120) else LIGHT_TEAL,
                    TEAL if x in (620,1620) else BLUE)
        draw.rounded_rectangle((x, 300, x + 360, 390), radius=24, fill=TEAL if x in (620,1620) else BLUE)
        centered_multiline(draw, (x + 20, 310, x + 340, 382), head, 30, "#FFFFFF", True)
        centered_multiline(draw, (x + 24, 415, x + 336, 635), body, 29, TEXT)
    for x in [480, 980, 1480]:
        arrow(draw, (x, 482), (x + 110, 482), color=ORANGE)
    rounded_box(draw, (240, 840, 1960, 1190), LIGHT_GRAY, NAVY, 30, 3)
    centered_multiline(draw, (280, 870, 1920, 940), "核心科学问题：如何在样本规模有限、钻孔方向不一致和缺陷尺度微小的条件下，稳定提取头部区域的三维判别特征？", 34, NAVY, True)
    modules = [
        (330, "避免错误裁剪", "TOML—RAW 索引对齐"),
        (790, "避免方向混淆", "统一头部相对位置"),
        (1250, "避免虚高结果", "板级锁定外部测试"),
    ]
    for x, h, b in modules:
        rounded_box(draw, (x, 1000, x + 350, 1135), "#FFFFFF", BLUE, 18, 2)
        centered_multiline(draw, (x + 10, 1008, x + 340, 1058), h, 27, NAVY, True)
        centered_multiline(draw, (x + 10, 1060, x + 340, 1125), b, 24, MID_GRAY)
    draw.text((100, 1328), "图 1  项目总体技术路线（自绘）", font=font(26), fill=MID_GRAY)
    img.save(path, quality=95)


def draw_network(path: Path):
    img, draw = canvas("Head40 三维 ResNet18 特征提取结构", "仅使用头部体积；首层沿钻孔长度方向采用 7×3×3 感受野", 1280)
    y1, y2 = 310, 715
    blocks = [
        (70, 300, "输入", "1×40×H×W\n方向已统一", LIGHT_BLUE),
        (400, 690, "Stem", "Conv 7×3×3\nBN + ReLU\nMaxPool 3×3×3, s=1", LIGHT_TEAL),
        (790, 1090, "Layer1", "2× BasicBlock\n64 通道\n保持分辨率", LIGHT_BLUE),
        (1190, 1480, "Layer2", "128 通道\nstride=2", LIGHT_TEAL),
        (1580, 1870, "Layer3/4", "256→512 通道\n分级下采样", LIGHT_BLUE),
    ]
    for x0, x1, h, b, fill in blocks:
        rounded_box(draw, (x0, y1, x1, y2), fill, TEAL if fill == LIGHT_TEAL else BLUE, 25, 3)
        centered_multiline(draw, (x0+15, y1+25, x1-15, y1+105), h, 33, NAVY, True)
        centered_multiline(draw, (x0+15, y1+125, x1-15, y2-25), b, 28, TEXT)
    for x in [310, 700, 1100, 1490]:
        arrow(draw, (x, 513), (x+75, 513), color=ORANGE)
    rounded_box(draw, (430, 850, 1770, 1115), "#F7FBFD", NAVY, 28, 3)
    centered_multiline(draw, (470, 875, 1730, 950), "全局聚合与分类：AdaptiveAvgPool → Dropout(0.5) → FC(2)", 35, NAVY, True)
    centered_multiline(draw, (470, 980, 1730, 1085), "最终特征图：29×29 输入横截面为 512×5×4×4；37×37 为 512×5×5×5；49×49 为 512×5×7×7。", 29, TEXT)
    draw.text((100, 1200), "图 2  Head40 三维 ResNet18 网络结构（自绘）", font=font(26), fill=MID_GRAY)
    img.save(path, quality=95)


def draw_validation(path: Path):
    img, draw = canvas("统一实验验证与应用输出流程", "用固定五折与锁定外部板，评价模型的稳定性、泛化性和可部署性", 1360)
    steps = [
        (100, "数据冻结", "固定 22 板五折\nSampleID / SHA256 审计"),
        (570, "训练验证", "五折 OOF 概率\nbest_loss 锁定模型"),
        (1040, "外部测试", "3 块锁定板\n五折概率集成"),
        (1510, "应用输出", "样本级判别\n板级指标与报告"),
    ]
    for i, (x, h, b) in enumerate(steps):
        fill = LIGHT_BLUE if i % 2 == 0 else LIGHT_TEAL
        rounded_box(draw, (x, 310, x+360, 635), fill, BLUE if i % 2 == 0 else TEAL, 26, 3)
        centered_multiline(draw, (x+15, 335, x+345, 410), h, 32, NAVY, True)
        centered_multiline(draw, (x+15, 440, x+345, 605), b, 29, TEXT)
        if i < 3:
            arrow(draw, (x+360, 472), (x+445, 472), color=ORANGE)
    rounded_box(draw, (150, 820, 2050, 1160), LIGHT_GRAY, NAVY, 30, 3)
    centered_multiline(draw, (200, 845, 2000, 910), "评价维度与后续消融", 36, NAVY, True)
    metrics = [
        (245, "分类性能", "Loss / Accuracy / Precision / Recall / F1 / Specificity"),
        (750, "泛化性能", "AUC、AP、逐板结果、五折均值±标准差"),
        (1255, "结构选择", "ResNet18/34 等骨干对比；Head24/32/40/48/64 长度消融"),
    ]
    for x, h, b in metrics:
        rounded_box(draw, (x, 955, x+430, 1115), "#FFFFFF", BLUE, 18, 2)
        centered_multiline(draw, (x+15, 965, x+415, 1018), h, 27, NAVY, True)
        centered_multiline(draw, (x+15, 1025, x+415, 1105), b, 22, MID_GRAY)
    draw.text((100, 1278), "图 3  统一验证、模型选择与工程化输出流程（自绘）", font=font(26), fill=MID_GRAY)
    img.save(path, quality=95)


def set_run_font(run, size=12, bold=False, color=None, name="宋体"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    rfonts.set(qn("w:eastAsia"), "仿宋_GB2312" if not bold else "黑体")


def clear_paragraph(paragraph):
    p = paragraph._element
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def remove_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=12):
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(22)
    fmt.space_after = Pt(0)
    fmt.space_before = Pt(0)
    if indent:
        fmt.first_line_indent = Pt(24)


def reset_cell(cell, keep_label=False):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    paras = cell.paragraphs
    if keep_label:
        for p in paras[1:]:
            remove_paragraph(p)
        label = paras[0]
        style_paragraph(label, WD_ALIGN_PARAGRAPH.LEFT, False, 12)
        for r in label.runs:
            set_run_font(r, 12, True)
        return label
    for p in paras[1:]:
        remove_paragraph(p)
    p = paras[0]
    clear_paragraph(p)
    return p


def add_text(cell, text, kind="body"):
    p = cell.add_paragraph()
    if kind == "heading":
        style_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT, False)
        r = p.add_run(text)
        set_run_font(r, 12, True, (23, 74, 114))
    elif kind == "bullet":
        style_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT, False)
        p.paragraph_format.left_indent = Pt(24)
        r = p.add_run(text)
        set_run_font(r, 12, False)
    elif kind == "caption":
        style_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, False, 10.5)
        p.paragraph_format.line_spacing = Pt(18)
        r = p.add_run(text)
        set_run_font(r, 10.5, False, (90, 108, 120))
    else:
        style_paragraph(p)
        r = p.add_run(text)
        set_run_font(r, 12, False)
    return p


def add_figure(cell, image_path: Path, caption: str, description: str):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    shape = run.add_picture(str(image_path), width=Cm(15.2))
    shape._inline.docPr.set("descr", description)
    shape._inline.docPr.set("title", caption)
    add_text(cell, caption, "caption")


def set_short(cell, text, center=True, bold=False, size=11):
    p = reset_cell(cell)
    style_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT, False, size)
    r = p.add_run(text)
    set_run_font(r, size, bold)


def replace_blank_paragraphs_with_text(cell, paragraphs):
    reset_cell(cell, keep_label=True)
    for kind, text in paragraphs:
        add_text(cell, text, kind)


def fill_cover_and_basic(doc: Document):
    cover = doc.tables[0]
    set_short(cover.cell(0, 1), TITLE, center=True, bold=True, size=12)
    set_short(cover.cell(1, 1), "李小嘉", center=True, size=12)
    set_short(cover.cell(2, 1), "待填写", center=True, size=12)
    set_short(cover.cell(3, 1), "待填写", center=True, size=12)
    set_short(cover.cell(5, 1), "待填写", center=True, size=12)
    set_short(cover.cell(6, 1), "待填写", center=True, size=12)
    set_short(cover.cell(7, 1), "待填写", center=True, size=12)

    basic = doc.tables[1]
    set_short(basic.cell(0, 2), TITLE, center=False, bold=True, size=10.5)
    set_short(basic.cell(1, 2), "智能制造与机器视觉相关科研平台（请按实际依托平台名称填写）", center=False, size=10.5)
    set_short(basic.cell(2, 2), "李小嘉", size=10.5)
    set_short(basic.cell(2, 4), "待填写", size=10.5)
    set_short(basic.cell(2, 6), "待填写", size=10.5)
    set_short(basic.cell(2, 9), "待填写", size=10.5)
    set_short(basic.cell(3, 2), "待填写", size=10.5)
    set_short(basic.cell(3, 8), "待填写", size=10.5)
    set_short(basic.cell(4, 8), "待填写", size=10.5)
    set_short(basic.cell(5, 2), "待填写", size=10.5)
    set_short(basic.cell(5, 8), "待填写", size=10.5)
    for col in (2, 4, 6, 9):
        set_short(basic.cell(6, col), "待填写", size=10.5)
    set_short(basic.cell(7, 2), "待填写", center=False, size=10.5)
    set_short(basic.cell(8, 1), "待填写", center=False, size=10)
    set_short(basic.cell(9, 1), "待填写", center=False, size=10)
    set_short(basic.cell(10, 1), "待填写", center=False, size=10)
    set_short(basic.cell(11, 1), "待填写", center=False, size=10)
    set_short(basic.cell(12, 8), "待填写", size=10)
    set_short(
        basic.cell(13, 1),
        "已围绕钻孔三维体积缺陷分类完成数据审计、固定五折划分、TOML 头部裁剪、方向统一、Head40 三维 ResNet18 建模及锁定外部板验证等工作；论文、项目与获奖情况请按实际补充。",
        center=False,
        size=10,
    )
    set_short(basic.cell(14, 1), "无（如有奖惩情况请按实际补充）", center=False, size=10)


def fill_design(doc: Document, figures: dict[str, Path]):
    design = doc.tables[2]
    c1, c2, c3, c4, c5, c6 = [design.cell(i, 0) for i in range(6)]

    reset_cell(c1, keep_label=True)
    for kind, text in [
        ("heading", "1. 研究背景与实践意义"),
        ("body", "钻孔构件广泛应用于机械加工、装备制造和工程连接等场景。钻孔内部的孔壁划伤、崩边、异物、局部欠加工和微小结构异常，往往具有尺寸小、位置隐蔽、灰度差异弱等特点；若不能在生产环节及时识别，可能影响连接可靠性、装配精度和后续服役安全。传统人工目检或逐层浏览三维切片的方式依赖经验，效率低且一致性有限。随着工业 CT、三维成像和自动化检测设备的普及，直接利用钻孔三维体积数据开展智能判别，具有明显的工程需求和实践价值。"),
        ("body", "本项目面向真实钻孔三维体积，研究微小缺陷的二分类方法：正常样本记为 0，缺陷样本记为 1。目标不是简单追求单一数据集上的高准确率，而是建立从元数据核验、头部区域确定、方向标准化、三维网络训练到外部板测试的完整流程，使分类结果能够服务于实际质检中的快速筛查和复核排序。该工作有望减少人工逐层判读负担，为工业三维无损检测提供可复用的算法组件。"),
        ("heading", "2. 国内外研究现状与问题"),
        ("body", "深度学习已经成为视觉检测的重要技术路线。二维卷积网络在表面缺陷识别中应用成熟，但对于钻孔内部结构，单张切片难以同时保留沿孔深方向的连续形态和横截面局部细节。三维卷积可在深度、宽度和高度三个维度联合学习特征，残差网络通过跳跃连接缓解深层网络训练困难，因此适合构建体积数据分类基线。近年来，DenseNet、EfficientNet、MobileNet 等结构也为性能与效率折中提供了选择。"),
        ("body", "然而，直接将通用三维网络用于工业钻孔数据仍面临三类困难。其一，缺陷并非均匀分布于整段体积，头部区域常包含更具判别力的加工与几何信息；若完整体积中包含大量无关尾部层，有限样本下模型易被冗余背景干扰。其二，同一物理区域在不同采集方向下可能位于深度轴的相反端，若不统一方向，网络会把“位置翻转”误学成类别差异。其三，工业数据常以板为单位采集，随机切分容易造成同板样本泄漏，难以反映对新板件的泛化能力。"),
        ("body", "针对上述问题，本项目以 TOML 标注文件中的 center2、bHeadUp 等字段为先验，建立 TOML—RAW 序列映射和头部区域裁剪规则；随后在运行时统一头部相对方向，并用固定五折和锁定三块外部板进行验证。现有基线结果表明，采用头部区域的三维 ResNet18 相较完整体积在外部板总体 F1 上具有更优表现，为进一步开展骨干网络对比与裁剪长度消融提供了明确基础。"),
    ]:
        add_text(c1, text, kind)
    add_figure(c1, figures["route"], "图 1  钻孔三维缺陷智能分类总体技术路线", "钻孔三维缺陷分类的总体技术路线图")
    for kind, text in [
        ("heading", "3. 主要参考文献"),
        ("bullet", "[1] He K, Zhang X, Ren S, Sun J. Deep Residual Learning for Image Recognition. CVPR, 2016."),
        ("bullet", "[2] Tran D, Bourdev L, Fergus R, Torresani L, Paluri M. Learning Spatiotemporal Features with 3D Convolutional Networks. ICCV, 2015."),
        ("bullet", "[3] Çiçek Ö, Abdulkadir A, Lienkamp S S, Brox T, Ronneberger O. 3D U-Net: Learning Dense Volumetric Segmentation from Sparse Annotation. MICCAI, 2016."),
        ("bullet", "[4] Litjens G, Kooi T, Bejnordi B E, et al. A Survey on Deep Learning in Medical Image Analysis. Medical Image Analysis, 2017, 42: 60-88."),
        ("bullet", "[5] Shorten C, Khoshgoftaar T M. A Survey on Image Data Augmentation for Deep Learning. Journal of Big Data, 2019, 6: 60."),
        ("bullet", "[6] Huang G, Liu Z, van der Maaten L, Weinberger K Q. Densely Connected Convolutional Networks. CVPR, 2017."),
        ("bullet", "[7] Tan M, Le Q. EfficientNet: Rethinking Model Scaling for Convolutional Neural Networks. ICML, 2019."),
        ("bullet", "[8] Howard A G, Zhu M, Chen B, et al. MobileNets: Efficient Convolutional Neural Networks for Mobile Vision Applications. arXiv:1704.04861, 2017."),
        ("bullet", "[9] Selvaraju R R, Cogswell M, Das A, et al. Grad-CAM: Visual Explanations from Deep Networks via Gradient-based Localization. ICCV, 2017."),
        ("bullet", "[10] Goodfellow I, Bengio Y, Courville A. Deep Learning. MIT Press, 2016."),
    ]:
        add_text(c1, text, kind)

    reset_cell(c2, keep_label=True)
    for kind, text in [
        ("heading", "1. 研究内容"),
        ("body", "（1）构建面向钻孔三维体积的数据规范与方向统一流程。解析 RAW 体数据及其 TOML 元数据，建立 file_sequence 到实际层序列的可审计映射；以 center2 确定中心层，以 bHeadUp 决定裁剪区间及运行时是否沿深度轴翻转，保证所有输入样本的头部均位于统一的相对方向。对体素值采用样本级 P1/P99 鲁棒归一化，并保留原始物理数据不被改写。"),
        ("body", "（2）研究头部三维体积的表征与分类方法。以 Head40 三维 ResNet18 为主基线，首层采用 7×3×3 非对称感受野，在长度方向捕获连续层结构、在横截面方向保持局部细节；在固定训练流程下，对 ResNet18、ResNet34、DenseNet121、EfficientNet-B0、MobileNetV2 等骨干进行公平比较，选择性能与计算代价均合适的方案。"),
        ("body", "（3）开展裁剪长度与泛化能力消融。基于统一规则构造 Head24、Head32、Head40、Head48、Head64 等不同深度长度的头部样本，比较信息保留与背景冗余的平衡；在相同五折划分、相同优化策略和相同锁定外部测试集上，报告 best_loss 对应模型、OOF 指标、逐板指标及五折概率集成结果。"),
        ("heading", "2. 研究目标"),
        ("body", "形成一套可复现的钻孔三维缺陷智能分类流程和数据审计规范；获得在锁定外部板上具有稳定 Accuracy、F1、AUC 和 AP 的三维分类模型；明确头部裁剪、方向统一、首层卷积核及网络骨干对性能的贡献；输出可供质检人员使用的样本级判别结果、板级统计报告和实验记录。项目验收时，力争外部测试的五折集成 F1 不低于现有 Head40 基线水平，并保持良好的召回率与特异性平衡。"),
        ("heading", "3. 拟解决的关键问题"),
        ("bullet", "① 元数据到原始体数据的准确对齐问题：避免中心层映射偏移导致的错误裁剪。"),
        ("bullet", "② 钻孔头尾方向的一致表征问题：消除 bHeadUp 不同造成的空间位置混淆。"),
        ("bullet", "③ 微小缺陷与冗余体积背景的平衡问题：确定合适的头部深度范围和三维感受野。"),
        ("bullet", "④ 有限工业样本下的可信验证问题：防止同板或同源样本泄漏，真实评估对新板件的泛化能力。"),
    ]:
        add_text(c2, text, kind)

    reset_cell(c3, keep_label=True)
    for kind, text in [
        ("heading", "1. 总体方案"),
        ("body", "项目按照“数据可信—表征有效—评估严格—结果可用”的思路实施。首先，对 22 块板的训练验证数据和 3 块锁定外部板进行文件、尺寸、标签、SampleID 和 SHA256 审计；随后从 TOML 获取中心位置与头部朝向，在线生成头部体积并完成方向统一。训练阶段使用固定五折清单，保证每次网络、裁剪长度和卷积核实验仅改变一个变量。测试阶段使用每折 best_loss 权重，对锁定三板分别推理并做五折概率集成。"),
        ("heading", "2. 关键方法"),
        ("body", "（1）TOML 引导的头部裁剪。将 TOML 的 file_sequence 转为从 0 开始的层序号，采用 center_index=floor(center2+0.5) 定位中心层。对 bHeadUp=false 的样本，以 [center-15, center+25) 构造 Head40；对 bHeadUp=true 的样本，以 [center-25, center+15) 构造相同物理意义的区域。超出边界的部分采取可审计的边界处理，确保样本大小一致。"),
        ("body", "（2）方向标准化与鲁棒输入。仅在数据加载时对 bHeadUp=true 的裁剪体沿深度维翻转，使头部统一位于高深度索引方向；不改写磁盘上的 RAW 数据。对每个样本计算 P1/P99 并截断归一化到 [0,1]，减弱不同采集批次灰度尺度差异。增强策略仅使用不会破坏深度方向语义的轻量空间和灰度扰动，避免随机深度翻转。"),
    ]:
        add_text(c3, text, kind)
    add_figure(c3, figures["network"], "图 2  Head40 三维 ResNet18 网络结构", "Head40 三维 ResNet18 的网络结构图")
    for kind, text in [
        ("body", "（3）三维残差分类网络。网络输入为 1×40×H×W 的单通道体积，首层为 1→64 通道的 7×3×3 卷积，步长为 1；随后最大池化同样不在首层下采样，以尽量保留微小缺陷的浅层细节。四个残差 stage 的通道为 64→128→256→512，仅在 layer2、layer3、layer4 做分级下采样，最后经自适应全局池化、Dropout 和二分类全连接层输出结果。对于 29×29、37×37、49×49 横截面，最终特征图分别为 512×5×4×4、512×5×5×5、512×5×7×7。"),
        ("body", "（4）公平比较与统计分析。训练采用相同的五折清单、类别标签、归一化、数据增强、优化器和训练轮数；以验证损失最低的 best_loss 权重作为每折锁定模型。除 Accuracy 外，重点报告 Precision、Recall、F1、Specificity、AUC 和 AP，并给出混淆矩阵、ROC/PR 曲线、五折均值±标准差以及逐板结果。针对类不平衡，综合关注召回率与误报率，避免仅凭准确率做结论。"),
        ("heading", "3. 实验设计与可行性"),
        ("body", "实验分为四步：第一步，复核完整体积基线与 Head40 基线，确认头部区域和方向统一的收益；第二步，完成 3×3×3、5×3×3、7×3×3 首层卷积核比较，现有结果显示 7×3×3 在外部总体表现上略优；第三步，比较不同骨干网络，回答选择 ResNet18 的依据；第四步，执行 Head24/32/40/48/64 长度消融，确定最适合钻孔头部微小缺陷的空间范围。所有实验均不改变锁定外部三板，保证结论可比。"),
    ]:
        add_text(c3, text, kind)
    add_figure(c3, figures["validation"], "图 3  统一验证、模型选择与工程化输出流程", "固定五折与锁定外部板的模型验证流程图")
    for kind, text in [
        ("body", "项目具有良好可行性。数据层面，已具备带 TOML 标注的三维 RAW 体数据、固定五折划分和锁定外部板；算法层面，已完成独立的训练、测试、数据加载和结果输出脚本，能够保存模型元数据、清单校验和完整指标；计算层面，可利用现有 GPU 训练环境完成三维网络迭代。现阶段 22 板训练验证样本共 4364 个、锁定外部样本 778 个，足以支持不同结构的同协议比较。"),
    ]:
        add_text(c3, text, kind)

    reset_cell(c4, keep_label=True)
    for kind, text in [
        ("bullet", "① 将 TOML 工艺/标注信息显式纳入三维数据构建。不是对完整体积做盲目截断，而是用 center2 和 bHeadUp 形成可解释、可复核的头部 ROI 规则。"),
        ("bullet", "② 提出“裁剪区间—运行时方向统一”两阶段策略。既保持原始 RAW 物理数据不变，又使网络输入具有一致的头部空间语义，降低方向混杂带来的学习难度。"),
        ("bullet", "③ 面向微小缺陷设计保持浅层分辨率的三维残差基线。首层采用 7×3×3 非对称感受野、最大池化不下采样，使长度连续信息与横截面局部纹理能够协同表达。"),
        ("bullet", "④ 建立板级锁定外部验证规范。利用固定五折 OOF、独立外部三板、逐板指标和五折集成，区分训练内性能与真实泛化能力，提升结果的工程可信度。"),
    ]:
        add_text(c4, text, kind)

    reset_cell(c5, keep_label=True)
    for kind, text in [
        ("heading", "1. 总体安排"),
        ("body", "2026 年 9 月—10 月：完成数据文件、TOML 字段、标签与固定五折清单的复核，形成数据规范和头部裁剪审计报告。"),
        ("body", "2026 年 11 月—2027 年 1 月：完成完整体积与 Head40 基线复现，完成 3×3×3、5×3×3、7×3×3 首层卷积核对比，固化最优基线。"),
        ("body", "2027 年 2 月—4 月：完成 ResNet18、ResNet34、DenseNet121、EfficientNet-B0、MobileNetV2 等骨干网络的统一协议比较，分析参数量、速度和精度。"),
        ("body", "2027 年 5 月—6 月：完成 Head24、Head32、Head40、Head48、Head64 等裁剪长度消融，开展外部板逐板分析和可解释性可视化。"),
        ("body", "2027 年 7 月—8 月：整理源代码、模型配置、实验记录和应用报告，撰写论文/软件著作权材料，完成结题与成果归档。"),
        ("heading", "2. 预期成果"),
        ("bullet", "① 形成一套钻孔三维体积缺陷分类的数据处理与方向统一规范；"),
        ("bullet", "② 形成可复现实验代码、五折模型权重、测试报告及板级结果汇总；"),
        ("bullet", "③ 发表或完成与项目相关的学术论文/会议论文 1 篇以上（以实际投稿和录用情况为准）；"),
        ("bullet", "④ 视研究进展申请软件著作权或整理工程应用技术报告 1 项。"),
    ]:
        add_text(c5, text, kind)

    reset_cell(c6, keep_label=True)
    for kind, text in [
        ("heading", "1. 已有研究基础"),
        ("body", "申请人已围绕钻孔三维缺陷分类开展前期工作，完成了数据文件检查、样本标签统计、固定五折划分及训练/测试流程搭建。当前训练验证数据包含 22 块板、4364 个样本，其中正常 3181 个、缺陷 1183 个；锁定外部测试数据包含 3 块板、778 个样本，其中正常 556 个、缺陷 222 个。划分过程中对板号、SampleID 与 RAW 文件 SHA256 进行核验，降低同源泄漏风险。"),
        ("body", "在模型探索方面，已完成完整体积多板 ResNet18、Head40-3×3×3、Head40-5×3×3、Head40-7×3×3 等实验。现有结果显示，头部裁剪模型总体优于完整体积基线；在锁定外部三板的五折集成中，Head40-7×3×3 的 Accuracy 约为 94.09%、F1 约为 89.82%、AUC 约为 98.35%，证明该路线具有可继续优化的潜力。非对称卷积分支的初步尝试未带来稳定收益，后续将把资源集中于骨干网络和裁剪长度的规范消融。"),
        ("heading", "2. 现有条件与保障"),
        ("body", "项目已具备三维 RAW 数据、TOML 元数据、Python/PyTorch 训练环境和 GPU 计算条件；现有代码已实现模型、训练、测试与结果汇总的独立管理，可自动导出 CSV/XLSX/JSON、ROC/PR 曲线和混淆矩阵。后续将在导师指导和现有实验平台支持下，按固定方案推进数据复核、模型对比和工程验证。涉及具体平台名称、导师信息及经费条件的内容，请在提交前按实际情况补全。"),
    ]:
        add_text(c6, text, kind)

    # Budget rows are deliberately labelled as an adjustable 1.00 万元 sample.
    set_short(design.cell(8, 0), "计算与数据存储服务费", size=10.5)
    set_short(design.cell(8, 1), "0.40", size=10.5)
    set_short(design.cell(8, 2), "用于 GPU 计算、数据备份与实验结果存储（按学校财务规定据实调整）。", center=False, size=10)
    set_short(design.cell(9, 0), "软件、测试与材料费", size=10.5)
    set_short(design.cell(9, 1), "0.25", size=10.5)
    set_short(design.cell(9, 2), "用于数据整理、模型测试、文献检索和必要的实验耗材。", center=False, size=10)
    set_short(design.cell(10, 0), "调研、差旅与成果费", size=10.5)
    set_short(design.cell(10, 1), "0.35", size=10.5)
    set_short(design.cell(10, 2), "用于现场调研、学术交流、论文版面或成果整理等（以实际资助额度为准）。", center=False, size=10)
    set_short(design.cell(11, 1), "1.00", size=10.5)
    set_short(design.cell(11, 2), "示例预算，提交前请按实际获批额度和学校财务规定核定。", center=False, size=10)


def count_cjk(text: str) -> int:
    return sum("\u4e00" <= char <= "\u9fff" for char in text)


def verify_document(path: Path):
    doc = Document(path)
    if len(doc.tables) != 4:
        raise RuntimeError(f"Unexpected table count: {len(doc.tables)}")
    if TITLE not in "\n".join(p.text for p in doc.paragraphs + [pp for t in doc.tables for c in t._cells for pp in c.paragraphs]):
        raise RuntimeError("Project title is missing from output")
    rels = [rel.target_ref for rel in doc.part.rels.values() if "image" in rel.reltype]
    if len(rels) < 3:
        raise RuntimeError(f"Expected 3 embedded images, found {len(rels)}")
    all_text = "\n".join(p.text for t in doc.tables for c in t._cells for p in c.paragraphs)
    return {"tables": len(doc.tables), "images": len(rels), "cjk_chars": count_cjk(all_text), "bytes": path.stat().st_size}


def main():
    if not SOURCE_TEMPLATE.exists():
        raise FileNotFoundError(SOURCE_TEMPLATE)
    OUTPUT_DIR.mkdir(exist_ok=True)
    figures_dir = WORK / "figures"
    figures_dir.mkdir(exist_ok=True)
    figures = {
        "route": figures_dir / "technical_route.png",
        "network": figures_dir / "head40_network.png",
        "validation": figures_dir / "validation_flow.png",
    }
    draw_technical_route(figures["route"])
    draw_network(figures["network"])
    draw_validation(figures["validation"])

    original_hash = sha256(SOURCE_TEMPLATE)
    shutil.copy2(SOURCE_TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    fill_cover_and_basic(doc)
    fill_design(doc, figures)
    doc.save(OUTPUT)
    if sha256(SOURCE_TEMPLATE) != original_hash:
        raise RuntimeError("Source template changed unexpectedly")
    result = verify_document(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"SOURCE_TEMPLATE_SHA256={original_hash}")
    for key, value in result.items():
        print(f"{key.upper()}={value}")


if __name__ == "__main__":
    main()
