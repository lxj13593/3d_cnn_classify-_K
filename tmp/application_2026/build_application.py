from __future__ import annotations

import hashlib
import math
from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\pythonproject\3d_cnn_classify _K")
TASK_DIR = ROOT / "tmp" / "application_2026"
SOURCE = TASK_DIR / "source.docx"
OUT_DIR = ROOT / "output"
FINAL = OUT_DIR / "李小嘉2026年山西省研究生实践创新项目申请书_钻孔三维缺陷分类.docx"
FIG_DIR = TASK_DIR / "figures"
EXPECTED_SHA256 = "363EC144B0AC367DF42DDE9E04F81A8567B5AC174BD7E9D57E062A911487C124"


PROJECT_TITLE = "基于方向统一与头部体积表征的钻孔三维缺陷智能分类方法研究"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest().upper()


def chinese_font_path(bold: bool = False) -> Path:
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc") if bold else Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\msyh.ttf"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError("No usable Chinese font found")


def pil_font(size: int, bold: bool = False):
    return ImageFont.truetype(str(chinese_font_path(bold=bold)), size=size)


def multiline_bbox(draw: ImageDraw.ImageDraw, text: str, font, spacing=8):
    return draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")


def draw_box(draw, xy, text, fill, font_size=31, bold=False, outline="#38516B", radius=22):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=4)
    font = pil_font(font_size, bold=bold)
    bb = multiline_bbox(draw, text, font, spacing=8)
    tw, th = bb[2] - bb[0], bb[3] - bb[1]
    draw.multiline_text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2), text, font=font, fill="#17202A", spacing=8, align="center")


def draw_arrow(draw, start, end, color="#607D8B", width=6, head=20):
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=color, width=width)
    angle = math.atan2(y2 - y1, x2 - x1)
    left = (x2 - head * math.cos(angle - math.pi / 6), y2 - head * math.sin(angle - math.pi / 6))
    right = (x2 - head * math.cos(angle + math.pi / 6), y2 - head * math.sin(angle + math.pi / 6))
    draw.polygon([(x2, y2), left, right], fill=color)


def centered_text(draw, xy, text, size, bold=False, color="#183A5A"):
    font = pil_font(size, bold=bold)
    bb = multiline_bbox(draw, text, font, spacing=8)
    tw, th = bb[2] - bb[0], bb[3] - bb[1]
    x, y = xy
    draw.multiline_text((x - tw / 2, y - th / 2), text, font=font, fill=color, spacing=8, align="center")


def make_technical_route(path: Path):
    image = Image.new("RGB", (2800, 1500), "white")
    draw = ImageDraw.Draw(image)
    centered_text(draw, (1400, 105), "钻孔三维缺陷分类总体技术路线", 58, bold=True)
    xs = [80, 620, 1160, 1700, 2240]
    w, h = 390, 245
    top = [
        ("原始三维体数据\n板号/孔号/标签清单", "#EAF2F8"),
        ("TOML 定位\ncenter2 + bHeadUp", "#D6EAF8"),
        ("方向统一\n钻头位于 D 轴高索引侧", "#D5F5E3"),
        ("头部 ROI\nHead24/32/40/48/64", "#FCF3CF"),
        ("逐样本 P1/P99\n按形状组批", "#FDEBD0"),
    ]
    for x, (text, color) in zip(xs, top):
        draw_box(draw, (x, 260, x + w, 260 + h), text, color, 30)
    for i in range(4):
        draw_arrow(draw, (xs[i] + w + 10, 382), (xs[i + 1] - 10, 382))
    bottom = [
        ("统一协议主干比较\n3D ResNet18/34\nDenseNet/EfficientNet/MobileNet", "#E8DAEF"),
        ("头部专用网络\n7×3×3 Stem\n浅层不下采样", "#D6EAF8"),
        ("22 板五折 OOF\nbest_loss\n多指标评估", "#D5F5E3"),
        ("锁定 3 块外部板\n778 个样本\n不参与模型选择", "#F9E79F"),
        ("3D 可解释分析\n错误归因与原型验证", "#F5CBA7"),
    ]
    for x, (text, color) in zip(xs, bottom):
        draw_box(draw, (x, 740, x + w, 1040), text, color, 28)
    for i in range(4):
        draw_arrow(draw, (xs[i] + w + 10, 890), (xs[i + 1] - 10, 890))
    draw_arrow(draw, (xs[-1] + w / 2, 515), (xs[-1] + w / 2, 730))
    draw.rounded_rectangle((150, 1190, 2650, 1375), radius=28, fill="#F4F6F7", outline="#AAB7B8", width=4)
    centered_text(
        draw,
        (1400, 1282),
        "数据层：物理方向一致、样本可追溯、训练测试板级隔离    |    模型层：全局与局部信息平衡\n评价层：验证与外部泛化分离",
        30,
        color="#37474F",
    )
    image.save(path, quality=95)


def make_model_architecture(path: Path):
    image = Image.new("RGB", (3000, 1500), "white")
    draw = ImageDraw.Draw(image)
    centered_text(draw, (1500, 105), "方向统一 Head40-733 三维残差网络结构", 58, bold=True)
    xs = [50, 465, 880, 1295, 1710, 2125, 2540]
    items = [
        ("输入\n1×40×H×W\nH/W=29,37,49", "#EAF2F8"),
        ("Stem Conv3D\n7×3×3, s=1\n1→64", "#D6EAF8"),
        ("MaxPool3D\n3×3×3, s=1\n不下采样", "#D5F5E3"),
        ("Layer1\nBasicBlock×2\n64, s=1", "#FCF3CF"),
        ("Layer2\nBasicBlock×2\n128, 首块 s=2", "#FDEBD0"),
        ("Layer3\nBasicBlock×2\n256, 首块 s=2", "#F5CBA7"),
        ("Layer4\nBasicBlock×2\n512, 首块 s=2", "#E8DAEF"),
    ]
    w = 330
    for x, (text, color) in zip(xs, items):
        draw_box(draw, (x, 300, x + w, 640), text, color, 28)
    for i in range(6):
        draw_arrow(draw, (xs[i] + w + 8, 470), (xs[i + 1] - 8, 470))
    draw_box(draw, (1250, 850, 1780, 1080), "AdaptiveAvgPool3D(1)\nDropout=0.5", "#D6EAF8", 34)
    draw_box(draw, (2080, 850, 2500, 1080), "Linear 512→2\n正常 / 缺陷", "#D5F5E3", 34, bold=True)
    draw_arrow(draw, (2705, 650), (1760, 840))
    draw_arrow(draw, (1795, 965), (2065, 965))
    draw.rounded_rectangle((120, 790, 980, 1130), radius=28, fill="#F8F9F9", outline="#95A5A6", width=4)
    centered_text(draw, (550, 960), "Layer4 最终特征图\n29×29 → 512×5×4×4\n37×37 → 512×5×5×5\n49×49 → 512×5×7×7", 31, color="#37474F")
    centered_text(draw, (1500, 1310), "核心设计：短深度头部体积浅层不下采样；7×3×3 同时建模长度方向连续性与横截面局部纹理", 34, color="#37474F")
    image.save(path, quality=95)


def make_results_plot(path: Path):
    image = Image.new("RGB", (2600, 1450), "white")
    draw = ImageDraw.Draw(image)
    centered_text(draw, (1300, 90), "前期实验结果对比（统一使用 best_loss 检查点）", 54, bold=True)
    models = ["FullVolume", "Head40-333", "Head40-533", "Head40-733", "Head40-733+ACB"]
    val_f1 = [87.64, 88.31, 88.88, 88.86, 88.58]
    test_f1 = [86.50, 88.13, 88.26, 88.58, 88.50]
    x0, y0, x1, y1 = 230, 210, 2460, 1150
    draw.line((x0, y0, x0, y1), fill="#455A64", width=4)
    draw.line((x0, y1, x1, y1), fill="#455A64", width=4)
    ymin, ymax = 84.0, 90.0
    for tick in range(84, 91):
        y = y1 - (tick - ymin) / (ymax - ymin) * (y1 - y0)
        draw.line((x0, y, x1, y), fill="#D5DBDB", width=2)
        centered_text(draw, (160, y), str(tick), 25, color="#455A64")
    centered_text(draw, (85, 675), "F1（%）", 30, color="#455A64")
    group_w = (x1 - x0) / len(models)
    bar_w = 105
    colors = ["#5DADE2", "#58D68D"]
    for i, (name, v1, v2) in enumerate(zip(models, val_f1, test_f1)):
        cx = x0 + group_w * (i + 0.5)
        for j, value in enumerate((v1, v2)):
            bx = cx + (-bar_w - 8 if j == 0 else 8)
            by = y1 - (value - ymin) / (ymax - ymin) * (y1 - y0)
            draw.rectangle((bx, by, bx + bar_w, y1), fill=colors[j], outline="#38516B", width=2)
            centered_text(draw, (bx + bar_w / 2, by - 24), f"{value:.2f}", 23, color="#263238")
        centered_text(draw, (cx, 1205), name, 24, color="#263238")
    draw.rectangle((280, 255, 335, 305), fill=colors[0])
    draw.text((350, 255), "五折验证 F1", font=pil_font(27), fill="#263238")
    draw.rectangle((280, 330, 335, 380), fill=colors[1])
    draw.text((350, 330), "锁定三板测试 F1", font=pil_font(27), fill="#263238")
    centered_text(draw, (1300, 1360), "外部测试为五折模型在三块锁定板上的均值；测试集不参与阈值、结构或检查点选择。", 27, color="#455A64")
    image.save(path, quality=95)


def set_run_font(run, size=10.5, bold=False, name="宋体", color=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def clear_cell(cell):
    tc = cell._tc
    for child in list(tc):
        if child.tag == qn("w:p"):
            tc.remove(child)


def allow_row_break(row, remove_height=False):
    trPr = row._tr.get_or_add_trPr()
    for node in list(trPr):
        if node.tag == qn("w:cantSplit"):
            trPr.remove(node)
        if remove_height and node.tag == qn("w:trHeight"):
            trPr.remove(node)


def add_para(cell, text="", kind="body", before=0, after=0):
    p = cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.widow_control = True
    if kind == "section":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.keep_with_next = True
        pf.line_spacing = 1.35
        r = p.add_run(text)
        set_run_font(r, size=11, bold=True, name="黑体")
    elif kind == "heading":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.keep_with_next = True
        pf.line_spacing = 1.35
        pf.space_before = Pt(max(before, 4))
        r = p.add_run(text)
        set_run_font(r, size=10.5, bold=True, name="黑体")
    elif kind == "reference":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.left_indent = Cm(0.63)
        pf.first_line_indent = Cm(-0.63)
        pf.line_spacing = 1.15
        r = p.add_run(text)
        set_run_font(r, size=10, name="宋体")
    elif kind == "note":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.left_indent = Cm(0.35)
        pf.right_indent = Cm(0.35)
        pf.line_spacing = 1.3
        r = p.add_run(text)
        set_run_font(r, size=9.5, name="楷体")
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(0.74)
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        r = p.add_run(text)
        set_run_font(r, size=10.5, name="宋体")
    return p


def add_picture(cell, path: Path, width_cm=14.6):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    return p


def add_caption(cell, text):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=9, name="宋体")
    return p


def set_placeholder_paragraph(p, text):
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=10.5, name="宋体")


def set_value_cell(cell, text, placeholder=False, center=True, size=10.5):
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size=size, name="宋体")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return p


def replace_cell_text_preserving_first_para(cell, text, placeholder=False):
    p = cell.paragraphs[0]
    if placeholder:
        set_placeholder_paragraph(p, text)
    else:
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=10.5, name="宋体")


def set_paragraph_bottom_border(p):
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")


def set_cell_edge(cell, edge, value="nil"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    node = borders.find(qn(f"w:{edge}"))
    if node is None:
        node = OxmlElement(f"w:{edge}")
        borders.append(node)
    node.set(qn("w:val"), value)


def clone_row_after(table, row_index):
    source_tr = table.rows[row_index]._tr
    source_tr.addnext(deepcopy(source_tr))


def unique_cells(row):
    cells = []
    seen = set()
    for c in row.cells:
        key = id(c._tc)
        if key not in seen:
            seen.add(key)
            cells.append(c)
    return cells


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ""
    run = p.add_run()
    set_run_font(run, size=9, name="宋体")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def fill_cover(doc: Document):
    table = doc.tables[0]
    values = [
        (PROJECT_TITLE, False),
        ("李小嘉", False),
        ("电子信息", False),
        ("0854", False),
        ("博士研究生□     硕士研究生☑", False),
        ("", False),
        ("中北大学                    （加盖公章）", False),
        ("2026 年 8 月 15 日", False),
    ]
    for row, (text, placeholder) in zip(table.rows, values):
        cells = unique_cells(row)
        target = cells[-1]
        p = set_value_cell(target, text, placeholder=placeholder, center=True, size=10.5)
        set_paragraph_bottom_border(p)


def fill_basic_data(doc: Document):
    table = doc.tables[1]
    # Project overview
    set_value_cell(unique_cells(table.rows[0])[-1], PROJECT_TITLE, center=True)
    set_value_cell(
        unique_cells(table.rows[1])[-1],
        "中北大学相关实验平台（正式名称待填写）",
        center=True,
        size=9.8,
    )

    # Applicant row: label/name, sex, birth, department.
    c = unique_cells(table.rows[2])
    if len(c) >= 9:
        set_value_cell(c[2], "李小嘉")
        set_value_cell(c[4], "")
        set_value_cell(c[6], "")
        set_value_cell(c[-1], "", size=9.5)

    c = unique_cells(table.rows[3])
    if len(c) >= 5:
        set_value_cell(c[2], "电子信息")
        set_value_cell(c[-1], "0854")

    c = unique_cells(table.rows[4])
    if len(c) >= 5:
        set_value_cell(c[2], "专业博士□  专业硕士☑")
        set_value_cell(c[-1], "是□     否☑")

    c = unique_cells(table.rows[5])
    if len(c) >= 5:
        set_value_cell(c[2], "")
        set_value_cell(c[-1], "是☑     否□")

    # Adviser fields.
    c = unique_cells(table.rows[6])
    if len(c) >= 9:
        set_value_cell(c[2], "", size=9.5)
        set_value_cell(c[4], "")
        set_value_cell(c[6], "")
        set_value_cell(c[-1], "")
    c = unique_cells(table.rows[7])
    if len(c) >= 3:
        set_value_cell(c[-1], "人工智能、机器视觉与工业无损检测", size=9.5)

    for row_idx, prompt in [
        (8, "待导师填写"),
        (9, "待导师填写"),
        (10, "待导师填写"),
        (11, "待导师填写"),
        (12, "是□     否□"),
    ]:
        c = unique_cells(table.rows[row_idx])
        set_value_cell(c[-1], prompt, center=False, size=9.5)

    # Applicant experience and awards.
    c = unique_cells(table.rows[13])
    narrative = c[-1]
    clear_cell(narrative)
    add_para(
        narrative,
        "已围绕钻孔三维体数据的智能分类开展前期研究，完成多板数据清理、样本追溯清单、正常/缺陷标签核对、孔级五折划分与外部板级隔离检查；已实现完整体积与头部体积的三维 ResNet18 训练、验证和独立测试流程，并形成可复现实验脚本与阶段性结果。相关成果目前处于研究与整理阶段。",
        kind="body",
    )
    c = unique_cells(table.rows[14])
    award = c[-1]
    clear_cell(award)
    add_para(award, "无。", kind="body")


def format_standalone_project_heading(doc: Document):
    for p in doc.paragraphs:
        if "项目设计论证" in p.text:
            p.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run("二、项目设计论证（建议5000字以内）")
            set_run_font(r, size=14, bold=False, name="宋体")
            return


SECTION_1_PARAS = [
    ("heading", "1. 研究背景与实践意义"),
    (
        "body",
        "钻孔是装备制造、电子封装、复合材料连接及精密构件加工中的基础工序。孔壁损伤、局部缺失、裂纹、毛刺、崩边或形貌异常会影响装配精度、连接强度和产品可靠性。传统人工复核主要依赖二维切片或特征图，面对批量生产中的大量孔位时，容易受到观察角度、人员经验和缺陷尺度的影响；单张二维图又难以完整表达缺陷沿孔深方向的连续形态。因此，面向真实生产数据建立可追溯、可复现、能够利用三维上下文的自动分类方法，具有降低复核成本、提高缺陷检出一致性和支撑制造过程闭环优化的实践价值。",
    ),
    (
        "body",
        "本项目以实际多板钻孔三维 RAW 体数据为对象，将每个孔表示为“长度 D × 高 H × 宽 W”的单通道体积。现有数据横截面包含 29×29、37×37 和 49×49 等规格，长度随扫描与裁剪设置变化。缺陷往往集中在钻头侧或孔口附近，但完整体积中同时包含大量正常孔段；若直接进行统一插值或过早下采样，可能削弱微小缺陷并引入形变。项目拟利用 TOML 中的 center2 与 bHeadUp 元数据完成头部定位和方向统一，在不修改原始 RAW 的前提下构建头部体积，再通过面向短深度输入的三维残差网络学习缺陷特征。",
    ),
    (
        "body",
        "从工程实施角度看，本研究不是单纯追求单一准确率，而是建立“数据治理—物理对齐—模型设计—严格验证—结果解释”的完整流程。训练/验证池由 22 块板、4364 个孔组成，其中正常 3181 个、缺陷 1183 个；另有 3 块从未参与训练与验证的外部板，共 778 个孔。训练与外部测试之间已经完成板号、SampleID 和 RAW SHA-256 去重检查。该数据基础能够支持同分布孔级泛化和未见板级泛化的分离评价，也为后续落地形成了真实场景验证条件。",
    ),
    ("heading", "2. 国内外研究现状"),
    (
        "body",
        "三维卷积网络能够直接在体素邻域内聚合长度、横向和纵向信息。Wu 等提出 3D ShapeNets，较早验证了体素表示在三维识别与形状补全中的可行性[1]；Maturana 和 Scherer 提出的 VoxNet 进一步展示了三维卷积在实时三维对象识别中的应用潜力[2]。残差学习通过恒等映射改善深层网络优化[3]，Hara 等系统研究了 3D ResNet 的深度与数据规模关系，说明三维残差结构能够形成有效的体积特征表征，但小样本条件下需注意过拟合[4]。DenseNet、MobileNetV2 和 EfficientNet 分别从特征复用、轻量化倒残差和复合尺度缩放角度提供了不同的主干设计[5-7]，适合作为同一协议下的候选网络。",
    ),
    (
        "body",
        "在工业无损检测领域，研究者已将深度学习用于工业 CT、层间监测、超声体数据和三维缺陷分析。Lee 等利用三维 CNN 对增材制造过程中的局部三维监测信号进行缺陷分类，并以 micro-CT 结果作为依据[8]；VolDefSegNet 面向纳米分辨率 X 射线 CT 的体缺陷分类与分割，表明体积上下文有助于区分形态相近的内部缺陷[9]。CTIMS 将 ResNet18 等架构用于工业 CT 工具完整性检查[10]。这些研究说明，三维表示能够避免逐切片方法割裂空间连续性，但工业数据仍普遍存在样本规模有限、设备与批次差异、标签成本高和解释性不足等问题。",
    ),
    (
        "body",
        "针对钻孔质量，已有工作从加工信号、表面多光照图像及 CAD 可制造性等角度开展研究。Sarkar 等采用 3D-CNN 识别 CAD 模型中难加工的局部钻孔特征，并将 Grad-CAM 扩展到三维对象[11]；Amini 等利用多光照成像与深度网络检测复合材料孔周损伤[12]；Schorr 等利用数控机床内部信号和机器学习预测钻孔与铰孔质量[13]。现有方法多针对二维孔口外观、过程信号或规则化 CAD 模型，对实际采集的可变尺寸钻孔三维体数据、钻头方向不一致和头部局部缺陷的联合处理研究仍然不足。",
    ),
    (
        "body",
        "综合来看，当前研究仍有三方面空缺：一是缺少利用采集元数据进行物理方向标准化的可复现流程；二是完整孔体积中的无关正常段可能稀释头部微缺陷，而通用 3D 网络的早期降采样又可能进一步损失细节；三是很多工作只在随机划分数据上报告准确率，未将未见板级外部测试与模型选择彻底隔离。本项目拟围绕这三点开展实践创新，并通过同一五折、同一外部测试和多指标体系验证方案。",
    ),
    ("heading", "3. 主要参考文献"),
]


REFERENCES = [
    "[1] Wu Z, Song S, Khosla A, et al. 3D ShapeNets: A deep representation for volumetric shapes[C]//CVPR. 2015: 1912-1920.",
    "[2] Maturana D, Scherer S. VoxNet: A 3D convolutional neural network for real-time object recognition[C]//IROS. 2015: 922-928.",
    "[3] He K, Zhang X, Ren S, Sun J. Deep residual learning for image recognition[C]//CVPR. 2016: 770-778.",
    "[4] Hara K, Kataoka H, Satoh Y. Can spatiotemporal 3D CNNs retrace the history of 2D CNNs and ImageNet?[C]//CVPR. 2018: 6546-6555.",
    "[5] Huang G, Liu Z, van der Maaten L, Weinberger K Q. Densely connected convolutional networks[C]//CVPR. 2017: 4700-4708.",
    "[6] Sandler M, Howard A, Zhu M, et al. MobileNetV2: Inverted residuals and linear bottlenecks[C]//CVPR. 2018: 4510-4520.",
    "[7] Tan M, Le Q V. EfficientNet: Rethinking model scaling for convolutional neural networks[C]//ICML. 2019: 6105-6114.",
    "[8] Lee K H, Lee H W, Yun G J. A defect detection framework using 3D-CNN with in-situ monitoring data in laser powder bed fusion process[J]. Optics & Laser Technology, 2023, 164: 109571.",
    "[9] Volumetric defect classification in nano-resolution X-ray computed tomography images of laser powder bed fusion via deep learning[J]. Journal of Manufacturing Processes, 2024, 121: 499-511. DOI: 10.1016/j.jmapro.2024.05.030.",
    "[10] CTIMS: Automated defect detection framework using computed tomography[J]. Applied Sciences, 2022, 12(4): 2175.",
    "[11] Sarkar S, et al. Learning localized features in 3D CAD models for manufacturability analysis of drilled holes[J]. Computer Aided Geometric Design, 2018, 62: 263-275.",
    "[12] Amini S, et al. Automated vision-based inspection of drilled CFRP composites using multi-light imaging and deep learning[J]. CIRP Journal of Manufacturing Science and Technology, 2021, 35: 441-453.",
    "[13] Schorr S, et al. In-process quality control of drilled and reamed bores using NC-internal signals and machine learning method[J]. Procedia CIRP, 2020, 93: 1328-1333.",
    "[14] Ding X, Guo Y, Ding G, Han J. ACNet: Strengthening the kernel skeletons for powerful CNN via asymmetric convolution blocks[C]//ICCV. 2019: 1911-1920.",
    "[15] Selvaraju R R, Cogswell M, Das A, et al. Grad-CAM: Visual explanations from deep networks via gradient-based localization[C]//ICCV. 2017: 618-626.",
    "[16] Loshchilov I, Hutter F. Decoupled weight decay regularization[C]//ICLR. 2019.",
]


SECTION_2_PARAS = [
    ("heading", "1. 研究内容"),
    (
        "body",
        "（1）建立面向多板钻孔体数据的数据治理与方向标准化流程。以板号、孔序号、标签、RAW 尺寸、TOML 序号和文件哈希为核心字段，形成可审计清单；依据 toml_sequence=file_sequence-1 对应元数据，采用 center_index=floor(center2+0.5) 定位头部中心，并根据 bHeadUp 决定是否沿 D 轴翻转，使所有样本的钻头侧统一位于高索引端。方向统一仅在内存加载阶段完成，不回写原始 RAW。",
    ),
    (
        "body",
        "（2）研究头部体积长度与有效信息范围。围绕 Head24、Head32、Head40、Head48 和 Head64 构建统一方向的头部 ROI，在保持 center2/bHeadUp 规则、五折清单、增强策略和训练参数一致的条件下，比较不同长度对验证损失、缺陷召回率、F1、AUC 与外部泛化的影响，分析过短裁剪造成信息不足和过长裁剪引入无关孔段之间的平衡。",
    ),
    (
        "body",
        "（3）开展统一协议下的三维主干网络比较。将 3D ResNet18、3D ResNet34、3D DenseNet121、3D EfficientNet-B0 和 3D MobileNetV2 应用于完整体积数据，固定数据划分、训练轮次、优化器、检查点规则和评价口径，比较精度、稳定性、参数量、显存占用及推理时间，为最终选择 ResNet18 或其他主干提供证据，避免只基于单一模型内部实验作结论。",
    ),
    (
        "body",
        "（4）构建头部体积专用三维残差分类网络。对 Head40 基线采用 7×3×3 非对称 Stem，长度方向感受野为 7、横截面保持 3×3；Stem 和首个最大池化均设 stride=1，避免在 40 层深度上过早压缩。Layer2、Layer3、Layer4 首块再执行 2 倍下采样，最后通过 AdaptiveAvgPool3d 接收不同横截面尺寸。进一步比较 Stem333、Stem533、Stem733，并将浅层 ACB 作为补充消融。",
    ),
    (
        "body",
        "（5）建立面向外部板泛化与解释性的评价体系。训练阶段报告五折 best_loss 检查点的 OOF 结果与均值±标准差；模型与阈值锁定后，仅在 3 块外部板上测试。除 Accuracy 外，同时报告 Precision、Recall、F1、Specificity、AUC、AP、混淆矩阵、逐板指标和逐尺寸指标；采用三维 Grad-CAM 与错误样本回溯观察网络关注区域是否位于钻头侧及缺陷邻域。",
    ),
    ("heading", "2. 研究目标"),
    (
        "body",
        "本项目拟形成一套适用于多板、可变尺寸钻孔体数据的可复现智能分类方案：完成方向标准化、头部裁剪、按形状组批、三维网络训练和锁定外部测试的完整工具链；在 22 板训练/验证池和 3 板外部测试集上形成无板号、SampleID 与 RAW 哈希泄露的规范评价；明确主干网络、头部长度和 Stem 长度的选择依据；力争在锁定外部测试上保持 Accuracy 不低于 93%、F1 不低于 88%、AUC 不低于 98%，并对误报与漏报来源给出可解释分析。最终形成可运行的软件原型、研究报告、学术论文或专利成果，为后续接入实际质检流程提供基础。",
    ),
    ("heading", "3. 拟解决的关键问题"),
    (
        "body",
        "（1）方向不一致导致同类结构表达相反的问题。必须利用 bHeadUp 提供的物理含义做确定性对齐，并禁止训练增强随机翻转 D 轴，否则网络会把方向差异误当成缺陷特征。",
    ),
    (
        "body",
        "（2）局部微缺陷与长体积冗余之间的矛盾。需要通过头部 ROI 长度消融确定有效范围，并调整 Stem 和池化步长，使浅层特征图保留足够的孔深连续性与横截面细节。",
    ),
    (
        "body",
        "（3）可变横截面尺寸与统一训练之间的矛盾。若强制插值到同一尺寸，可能改变缺陷尺度；本项目采用按 D×H×W 形状组批和自适应全局池化，既保持原始体素几何，又实现统一分类头。",
    ),
    (
        "body",
        "（4）小样本、类别不均衡与板间分布偏移问题。通过固定五折、轻量增强、AdamW 正则化、多个分类指标和板级外部测试控制过拟合，不根据外部测试反向选择结构、阈值或检查点。",
    ),
]


SECTION_3_INTRO = [
    ("heading", "1. 研究方法"),
    (
        "body",
        "本项目采用“数据审计—物理对齐—头部建模—受控消融—外部验证—解释与应用”的研究方法。所有实验均从同一份样本清单出发，以正常为 0、缺陷为 1；每个样本独立执行 P1/P99 截断归一化到[0,1]。训练增强仅作用于 H/W 平面与灰度，包括横截面翻转、90°旋转、±3 体素平移、高斯噪声、亮度和对比度轻微扰动；验证集和外部测试集不使用随机增强。",
    ),
]


SECTION_3_AFTER_ROUTE = [
    ("heading", "2. 技术路线与关键技术"),
    (
        "body",
        "（1）TOML 驱动的头部定位与方向统一。对于 bHeadUp=false 的样本，Head40 取[center_index-15, center_index+25)；对于 bHeadUp=true 的样本，取[center_index-25, center_index+15)。预裁剪体积仍保留原始方向，加载时仅对 bHeadUp=true 的样本沿 D 轴反转，从而使钻头端统一到高索引侧。该方法把设备元数据转化为网络可利用的物理先验，同时保持源数据不可变和全流程可追溯。",
    ),
    (
        "body",
        "（2）可变尺寸保持与按形状组批。三种横截面不进行插值，同尺寸样本组成一个 batch；网络末端使用自适应三维全局平均池化，将不同 D×H×W 特征统一为 512 维向量。该设计避免了人为缩放造成的尺度漂移，并可同时利用不同扫描规格的数据。",
    ),
    (
        "body",
        "（3）头部专用 3D ResNet。基础模型沿用[2,2,2,2]残差块和 64→128→256→512 通道配置，将完整体积模型 Stem 的深度下采样改为 stride=1，并取消首池化下采样。对 Head40-733，输入 40×29×29、40×37×37、40×49×49 时，Layer4 最终特征图分别为 512×5×4×4、512×5×5×5、512×5×7×7，仍保留可用于区分缺陷的空间单元。",
    ),
]


SECTION_3_AFTER_MODEL = [
    (
        "body",
        "（4）公平的模型与消融设计。第一阶段在完整体积上比较 ResNet18、ResNet34、DenseNet121、EfficientNet-B0 和 MobileNetV2；第二阶段固定主干，比较 FullVolume 与不同 Head 长度；第三阶段固定 Head 长度，比较 3×3×3、5×3×3 和 7×3×3 Stem；第四阶段以 ACB 为补充，验证方向性局部分支是否带来稳定收益。每组实验只改变一个核心变量，统一使用 50 epoch、batch size 4、AdamW、初始学习率 1e-4、weight decay 1e-3、余弦退火和未加权交叉熵。",
    ),
    (
        "body",
        "（5）检查点与评价规则。每折保存 best_loss、best_f1 和 last，论文主结果固定使用 best_loss，避免跨模型混用不同选择规则。五折验证同时报告 OOF 汇总损失与均值±标准差；外部测试报告五折模型在三块锁定板上的均值±标准差，概率集成仅作补充分析，不用于反向选择模型。",
    ),
    ("heading", "3. 实验手段"),
    (
        "body",
        "数据层面，训练/验证池包含 22 块板共 4364 个样本，外部测试包含 3 块板共 778 个样本。固定五折的验证样本数为 872、873、874、873、872，所有主要比较复用相同 SampleID 折归属。清单生成时核对 RAW 存在性、文件字节数、序号、标签、尺寸、bHeadUp、方向标准化状态及 SHA-256，保证样本可追溯。",
    ),
    (
        "body",
        "模型训练在 PyTorch 与 CUDA 环境中完成，启用自动混合精度以控制显存。对每个实验记录配置、随机种子、训练曲线、逐样本概率、混淆矩阵、ROC/PR 曲线和逐板指标；对错误样本回查原始体数据、二维特征图和元数据，区分标签边界模糊、局部低对比、板间分布差异及模型关注区域偏移等原因。",
    ),
    (
        "body",
        "模型比较除分类性能外，还统计参数量、FLOPs 或近似计算量、显存峰值与单样本推理时间，形成精度—效率双维度结论。对最终模型采用三维 Grad-CAM[15]生成体积热力图，并沿 D 轴和横截面投影，与钻头侧位置及人工复核缺陷区域进行对应，提升工程人员对模型输出的信任度。",
    ),
    ("heading", "4. 前期结果与后续验证"),
    (
        "body",
        "前期已完成 FullVolume、Head40-333、Head40-533、Head40-733 以及只修改 layer1[0].conv1 的 ACB 实验。Head40-733 相比 FullVolume，验证 OOF loss 由 0.167286 降至 0.152872；锁定三板五折平均 Accuracy 由 92.29% 提升至 93.37%，F1 由 86.50% 提升至 88.58%。Head40-533 的验证 loss 最低，Head40-733 的外部 Accuracy、Precision、F1、Specificity 和 AP 综合略优；ACB 未形成稳定增益。上述结果说明头部建模路线具有进一步研究价值，但不能替代待完成的主干比较和 Head24/32/48/64 长度消融。",
    ),
]


SECTION_3_AFTER_RESULTS = [
    ("heading", "5. 可行性分析"),
    (
        "body",
        "（1）数据可行性。现有 25 块板、5142 个样本已形成训练/验证与外部测试两套物理隔离目录，三种横截面和多种长度均有清单记录；样本与二维图一一对应，训练测试之间无板号、SampleID 或 RAW 哈希重复。",
    ),
    (
        "body",
        "（2）技术可行性。三维卷积、残差网络、自适应池化、AdamW 和 Grad-CAM 均有成熟理论与实现基础；项目现已完成数据加载、五折训练、检查点保存、逐板测试和图表输出，关键模块可在现有代码上按独立实验文件扩展。",
    ),
    (
        "body",
        "（3）实验可行性。现有 GPU 环境能够完成 batch size 4 的五折训练；相同五折清单与自动化结果汇总降低了重复实验的人为误差。后续主干比较只进行一轮统一协议试验，长度消融固定为五个候选值，工作量可在一年周期内完成。",
    ),
    (
        "body",
        "（4）应用可行性。输出为孔级正常/缺陷概率，可与板号、孔号和原始图关联，便于嵌入人工复核流程。通过外部板测试、逐板统计和三维热力图，可为生产现场评估漏检、误报和板间稳定性提供直接依据。",
    ),
]


SECTION_4_PARAS = [
    ("heading", "1. 项目特色"),
    (
        "body",
        "项目面向真实多板、可变尺寸钻孔体数据，不以理想化公共数据代替工程问题，将数据元信息、物理方向、三维局部结构与板级外部泛化统一到同一流程。研究既包含网络结构改进，也重视训练测试隔离、样本追溯和失败样本解释，具有较强的工程完整性。",
    ),
    ("heading", "2. 实践创新"),
    (
        "body",
        "（1）提出基于 center2 与 bHeadUp 的头部体积构建方法。与随机翻转或仅按几何中心裁剪不同，本项目利用采集元数据恢复钻头侧位置，使所有样本在物理方向上对齐，减少方向差异带来的伪特征。",
    ),
    (
        "body",
        "（2）提出兼顾长轴连续性与横截面局部纹理的 7×3×3 Stem，并在 Head40 浅层取消下采样。该设计不盲目扩大全部三维卷积核，而是在最接近原始缺陷的位置强化 D 方向上下文，后续仍使用标准残差块，保持结构清晰和可解释。",
    ),
    (
        "body",
        "（3）采用“不同横截面不插值 + 按形状组批 + 自适应池化”的可变尺寸处理方式。该方法保留原始体素尺度，避免插值平滑微小缺陷，同时允许同一模型处理 29×29、37×37 和 49×49 三种横截面。",
    ),
    (
        "body",
        "（4）建立验证集与外部板严格分离的评价协议。外部三板不参与五折、归一化统计、阈值选择、最佳轮次和结构选择；所有模型使用相同 SampleID 折归属和 best_loss 规则，避免因评价口径变化产生虚假提升。",
    ),
    (
        "body",
        "（5）形成从算法到质检复核的可追溯输出。每个预测均关联板号、孔号、概率、原始文件和三维关注区域，能够支持人工复核、误差分析和后续质量规则迭代，而不仅输出一个不可解释的分类标签。",
    ),
]


SECTION_5_PARAS = [
    ("heading", "1. 总体安排"),
    ("body", "（1）2026 年 9 月—10 月：复核个人与项目申报信息，冻结数据清单和统一实验协议；完成 FullVolume 下 ResNet18、ResNet34、DenseNet121、EfficientNet-B0、MobileNetV2 主干比较，统计性能与资源开销。"),
    ("body", "（2）2026 年 11 月—12 月：在选定主干上完成 Head24、Head32、Head40、Head48、Head64 长度消融，分析不同孔深范围对验证损失、召回率、F1 和外部稳定性的影响。"),
    ("body", "（3）2027 年 1 月—2 月：整合 Stem333/533/733 与 ACB 前期结果，确定最终模型；补充参数量、显存和推理时间测试，完成三维 Grad-CAM、错误样本回溯与逐板分析。"),
    ("body", "（4）2027 年 3 月—4 月：在锁定外部板和新增可用生产样本上开展一次性验证，整理模型校准、误报/漏报和不同板批差异；开发孔级批量推理与结果可视化原型。"),
    ("body", "（5）2027 年 5 月—6 月：完成实验表格、技术报告和论文初稿，凝练方向统一、头部 ROI 与三维结构设计的贡献，视成果成熟度申请软件著作权或专利。"),
    ("body", "（6）2027 年 7 月—8 月：完成成果修改、验收材料和应用演示，归档数据清单、独立脚本、模型权重与复现实验说明。"),
    ("heading", "2. 预期成果"),
    ("body", "（1）形成一套钻孔三维体数据清理、方向统一、头部裁剪、训练验证和外部测试的规范流程与软件原型。"),
    ("body", "（2）完成 5 类主干网络、5 种头部长度、3 种 Stem 长度及 ACB 补充消融，形成可复现的实验报告和最终模型。"),
    ("body", "（3）在 3 块锁定外部板上力争达到 Accuracy≥93%、F1≥88%、AUC≥98%，并给出逐板、逐尺寸和错误类型分析。"),
    ("body", "（4）发表或投稿学术论文 1 篇，完成软件著作权或专利申请 1 项，形成项目技术报告与演示材料。"),
    ("body", "（5）为后续与现场质检系统对接、增量扩展新板数据和开展缺陷细分类提供可持续的数据与代码基础。"),
]


SECTION_6_PARAS = [
    ("heading", "1. 过去的研究基础"),
    (
        "body",
        "申请人已围绕钻孔三维缺陷二分类完成较系统的前期工作。数据方面，已建立 22 板训练/验证池和 3 板锁定外部测试集，完成正常与缺陷标签整理、明显/模糊样本追溯、RAW 尺寸统计和训练测试泄露检查。当前训练/验证样本为 4364 个，外部测试样本为 778 个，所有数据均有板号、孔号、路径和哈希记录。",
    ),
    (
        "body",
        "算法方面，已实现完整体积 3D ResNet18、方向统一 Head40-333/533/733 以及浅层 ACB 模型，能够处理不同横截面尺寸；已完成固定五折训练、best_loss/best_f1/last 检查点管理、OOF 预测、逐板指标、ROC/PR 曲线和锁定三板测试。前期结果中，Head40-733 外部五折平均 Accuracy 为 93.37%±0.65%，F1 为 88.58%±0.91%，AUC 为 98.17%±0.17%，为项目进一步开展主干比较与长度消融提供了可靠起点。",
    ),
    (
        "body",
        "研究过程已形成“旧脚本不改、每个实验独立新建模型/训练/测试/说明文件”的管理原则，确保不同实验可追溯且互不覆盖。现有总结明确区分完整体积、头部长度、Stem 卷积核长度和 ACB 四类变量，并固定以 best_loss 为论文主结果，具备继续规范推进的研究基础。",
    ),
    ("heading", "2. 现有工作条件"),
    ("body", "（1）软件条件：已具备 Python、PyTorch、CUDA、NumPy、pandas、scikit-learn 和可视化工具链，可完成三维数据读取、混合精度训练、指标统计和图表输出。"),
    ("body", "（2）计算条件：现有 GPU 深度学习环境已验证可完成五折 3D ResNet18 训练；按形状组批和 batch size 4 能够在现有显存条件下稳定运行。具体设备型号及共享服务器信息请在提交前据实补充。"),
    ("body", "（3）数据条件：现有多板数据覆盖三种横截面和多种体积长度，训练/验证与外部测试在物理目录及清单层面隔离；数据来源、标签与文件完整性均可回查。"),
    ("body", "（4）指导与平台条件：项目可依托所在团队在人工智能、机器视觉和工业检测方面的研究积累开展。正式实验室名称、导师科研项目与经费情况需由申请人和导师在提交前核实填写。"),
    ("body", "（5）风险控制：针对过拟合、类别不均衡和板间偏移，项目采用固定五折、轻量增强、正则化、多指标和外部板验证；针对结果不可解释，采用三维热力图和错误样本回溯；针对实验量较大，采用分阶段单变量设计和自动化汇总。"),
    ("heading", "3. 实施保障与数据管理"),
    (
        "body",
        "项目实行清单化、版本化和分阶段验收管理。数据层保存板号、孔号、标签、RAW 路径、TOML 序号、体积尺寸、方向标记与文件哈希，训练前自动核验文件存在性、字节数及训练测试交集；实验层保存配置、随机种子、软件环境、训练日志、逐折检查点、OOF 概率和绘图源数据；成果层同步归档统计表、错误样本清单、可视化结果与结论说明。原始数据只读保存，派生清单和结果按实验名称独立建档，关键结果至少保留两份副本，从而保证项目过程可复查、结果可复现、失败实验可回溯。",
    ),
    (
        "body",
        "研究任务按“主干网络比较—头部长度消融—结构细化—锁定外部验证—解释与应用”五个阶段推进。每一阶段预先确定唯一变量、样本划分、检查点和评价口径，完成后再进入下一阶段；若出现外部性能明显下降、训练不稳定或资源消耗超出预期，则回到上一稳定基线分析原因，不根据外部测试结果反复选择模型。通过阶段评审、结果复核和统一汇总表控制进度与研究偏差。",
    ),
    ("heading", "4. 成果应用与条件完善"),
    (
        "body",
        "项目成果面向孔级质量复核场景组织输出。最终原型以单孔三维体及其元数据为输入，输出正常/缺陷概率、阈值判定、板号与孔号索引，并生成沿孔深方向和横截面的三维关注区域投影，便于检测人员回看原始数据。应用验证将关注漏检率、误报率、板间稳定性、单样本推理时间和结果可追溯性，同时记录不同尺寸、不同批次及边界样本上的失败模式。正式设备型号、平台名称、导师团队项目和配套条件将在申报材料定稿时由申请人与导师依据实际情况填写，不在研究内容中作未经核实的承诺。",
    ),
    ("heading", "5. 质量控制与阶段验收"),
    (
        "body",
        "项目设置数据、模型、指标和解释四级质量控制。数据验收要求训练/验证与锁定外部板无板号、SampleID 和 RAW 哈希交集，所有样本的标签、尺寸、方向及裁剪范围能够回查；模型验收要求网络结构、参数、训练轮次和检查点来源记录完整，同类实验仅改变预先声明的变量；指标验收要求五折验证统一报告均值与标准差，外部测试在模型和阈值锁定后执行，并同时给出 Accuracy、Precision、Recall、F1、Specificity、AUC、AP 与混淆矩阵；解释验收要求三维关注区域能够与钻头侧及缺陷邻域对应，典型误判样本形成原因分类。阶段成果由申请人整理原始记录、汇总表和复现实验说明，经导师复核后归档。对结论不稳定或不同批次差异较大的结果，不作选择性报告，而是补充逐板统计、置信区间和失败原因分析，保证项目结论真实、完整且可重复验证。",
    ),
]


def fill_project_design(doc: Document, route: Path, model_fig: Path, results_fig: Path):
    table = doc.tables[2]
    # Split the long first section into three visually continuous rows. This keeps
    # the first content on the heading page and avoids Word moving one huge row.
    clone_row_after(table, 0)
    clone_row_after(table, 1)
    for row in table.rows[:8]:
        allow_row_break(row, remove_height=True)
        c = unique_cells(row)[0]
        set_cell_margins(c, top=100, start=120, bottom=100, end=120)
        clear_cell(c)

    # Section 1, part A: background and significance.
    c = unique_cells(table.rows[0])[0]
    add_para(c, "（一）立项依据（包括研究背景与实践意义、国内外研究现状等，并附主要参考文献）", kind="section")
    for kind, text in SECTION_1_PARAS[:4]:
        add_para(c, text, kind=kind)
    set_cell_edge(c, "bottom", "nil")

    # Section 1, part B: research status. The suppressed horizontal borders make
    # the three rows appear as one continuous form field.
    c = unique_cells(table.rows[1])[0]
    for kind, text in SECTION_1_PARAS[4:9]:
        add_para(c, text, kind=kind)
    set_cell_edge(c, "top", "nil")
    set_cell_edge(c, "bottom", "nil")

    # Section 1, part C: references.
    c = unique_cells(table.rows[2])[0]
    for kind, text in SECTION_1_PARAS[9:]:
        add_para(c, text, kind=kind)
    for ref in REFERENCES:
        add_para(c, ref, kind="reference")
    set_cell_edge(c, "top", "nil")

    # Section 2
    c = unique_cells(table.rows[3])[0]
    add_para(c, "（二）研究内容、研究目标以及拟解决的关键问题（此部分为重点阐述内容）", kind="section")
    for kind, text in SECTION_2_PARAS:
        add_para(c, text, kind=kind)

    # Section 3
    c = unique_cells(table.rows[4])[0]
    add_para(c, "（三）拟采取的研究方案及其可行性分析（包括研究方法、技术路线、实验手段、关键技术等）", kind="section")
    for kind, text in SECTION_3_INTRO:
        add_para(c, text, kind=kind)
    add_picture(c, route, width_cm=14.8)
    add_caption(c, "图 1  钻孔三维缺陷分类总体技术路线")
    for kind, text in SECTION_3_AFTER_ROUTE:
        add_para(c, text, kind=kind)
    add_picture(c, model_fig, width_cm=14.8)
    add_caption(c, "图 2  方向统一 Head40-733 三维残差网络结构")
    for kind, text in SECTION_3_AFTER_MODEL:
        add_para(c, text, kind=kind)
    add_picture(c, results_fig, width_cm=14.5)
    add_caption(c, "图 3  前期五折验证与锁定三板测试 F1 对比")
    for kind, text in SECTION_3_AFTER_RESULTS:
        add_para(c, text, kind=kind)

    # Section 4
    c = unique_cells(table.rows[5])[0]
    add_para(c, "（四）本项目的特色及实践创新之处", kind="section")
    for kind, text in SECTION_4_PARAS:
        add_para(c, text, kind=kind)

    # Section 5
    c = unique_cells(table.rows[6])[0]
    add_para(c, "（五）研究工作的总体安排及预期成果", kind="section")
    for kind, text in SECTION_5_PARAS:
        add_para(c, text, kind=kind)

    # Section 6
    c = unique_cells(table.rows[7])[0]
    add_para(c, "（六）研究基础与工作条件（包括过去的研究基础，现有的研究条件等）", kind="section")
    for kind, text in SECTION_6_PARAS:
        add_para(c, text, kind=kind)

    # Budget table rows 8-13. Preserve heading and header row, fill 3 entries + total.
    row8 = unique_cells(table.rows[8])
    set_value_cell(row8[0], "（七）经费预算", center=False)
    entries = [
        ("计算与存储", "0.30", "用于 GPU 算力、数据备份存储及实验过程中必要的计算资源使用费用"),
        ("实验材料与数据整理", "0.10", "用于样本整理、移动存储介质、图像打印及现场验证所需的辅助材料"),
        ("成果与资料", "0.10", "用于文献资料、论文版面、软件著作权或专利申请及项目材料印制"),
    ]
    for row, (name, amount, reason) in zip(table.rows[10:13], entries):
        cells = unique_cells(row)
        set_value_cell(cells[0], name, center=True, size=9.5)
        set_value_cell(cells[1], amount, center=True, size=9.5)
        set_value_cell(cells[-1], reason, center=False, size=9.5)
    total = unique_cells(table.rows[13])
    set_value_cell(total[0], "合   计", center=True, size=9.5)
    set_value_cell(total[1], "0.50", center=True, size=9.5)
    set_value_cell(total[-1], "预算总额 0.50 万元，具体科目可按学校当年通知调整。", center=False, size=9.5)


def normalize_commitment_page(doc: Document):
    table = doc.tables[3]
    for row in table.rows:
        allow_row_break(row)
        for cell in unique_cells(row):
            set_cell_margins(cell, top=80, start=90, bottom=80, end=90)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.25
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)


def collapse_trailing_empty_paragraph(doc: Document):
    if not doc.paragraphs:
        return
    p = doc.paragraphs[-1]
    if p.text.strip():
        return
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(1)
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.widow_control = False
    r = p.add_run(" ")
    set_run_font(r, size=1, name="宋体")


def set_core_properties(doc: Document):
    doc.core_properties.title = PROJECT_TITLE
    doc.core_properties.subject = "2026年山西省研究生实践创新项目申请书"
    doc.core_properties.keywords = "钻孔；三维卷积神经网络；ResNet18；方向统一；头部体积；缺陷分类"


def main():
    if sha256(SOURCE) != EXPECTED_SHA256:
        raise RuntimeError("Retained DOCX hash mismatch; fresh distillation required")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    route = FIG_DIR / "technical_route.png"
    model_fig = FIG_DIR / "head40_resnet18_architecture.png"
    results_fig = FIG_DIR / "preliminary_results.png"
    make_technical_route(route)
    make_model_architecture(model_fig)
    make_results_plot(results_fig)

    doc = Document(str(SOURCE))
    fill_cover(doc)
    fill_basic_data(doc)
    format_standalone_project_heading(doc)
    fill_project_design(doc, route, model_fig, results_fig)
    collapse_trailing_empty_paragraph(doc)
    set_core_properties(doc)
    add_page_number(doc.sections[0])

    # Ensure Word refreshes PAGE fields on open.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(str(FINAL))
    print(FINAL)
    print(f"final_bytes={FINAL.stat().st_size}")
    print(f"source_sha256_after={sha256(SOURCE)}")


if __name__ == "__main__":
    main()
